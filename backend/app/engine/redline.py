"""Redline export: redlines.json schema + .docx builder.

python-docx exposes no tracked-changes API (`Document.paragraphs` skips
`<w:ins>`/`<w:del>` content), so deletions render as red-strikethrough runs
and insertions as red-underline runs — visually identical to Word redlines.
"""

import io
from typing import Literal

from pydantic import BaseModel, Field

try:
    from docx import Document
    from docx.shared import RGBColor

    _DOCX_OK = True
except ImportError:  # pragma: no cover
    _DOCX_OK = False

Action = Literal["insert", "delete", "modify"]


class RedlineCite(BaseModel):
    rule_id: str
    excerpt: str
    start: int
    end: int
    side: Literal["v1", "v2"]


class RedlineItem(BaseModel):
    clause: str
    action: Action
    v1_text: str = ""
    v2_text: str = ""
    cites: list[RedlineCite] = Field(default_factory=list)


class RedlinesJson(BaseModel):
    """`redlines.json` schema (legal-redline-tools pattern)."""

    version: str = "w2.2"
    items: list[RedlineItem] = Field(default_factory=list)


def to_redlines_json(deltas: list[dict]) -> RedlinesJson:
    items: list[RedlineItem] = []
    for d in deltas:
        if d.get("label") in ("cosmetic", "same"):
            continue
        action: Action = "modify"
        if d.get("label") == "added":
            action = "insert"
        elif d.get("label") == "deleted":
            action = "delete"
        cites = [RedlineCite(**c) for c in d.get("cites", [])]
        v1 = next((c.excerpt for c in cites if c.side == "v1"), "")
        v2 = next((c.excerpt for c in cites if c.side == "v2"), "")
        clause = cites[0].rule_id if cites else "general"
        items.append(
            RedlineItem(
                clause=clause, action=action, v1_text=v1, v2_text=v2, cites=cites
            )
        )
    return RedlinesJson(items=items)


def build_redline_docx(red: RedlinesJson) -> bytes:
    """Render redlines to .docx bytes; raises RuntimeError without python-docx."""
    if not _DOCX_OK:
        raise RuntimeError("python-docx is not installed")
    doc = Document()
    doc.add_heading("NyayaMitra — Contract Redline (v1 → v2)", level=1)
    for it in red.items:
        doc.add_heading(f"{it.clause} [{it.action}]", level=2)
        if it.v1_text:
            p = doc.add_paragraph()
            r = p.add_run(it.v1_text)
            r.font.strike = True
            r.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
        if it.v2_text:
            p = doc.add_paragraph()
            r = p.add_run(it.v2_text)
            r.font.underline = True
            r.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
